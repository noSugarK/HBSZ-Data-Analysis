import sys
import pandas as pd
import numpy as np
from PyQt6.QtWidgets import (QApplication, QMainWindow, QTabWidget, QWidget,
                             QVBoxLayout, QHBoxLayout, QLabel, QComboBox,
                             QPushButton, QFileDialog, QMessageBox, QGroupBox,
                             QTableWidget, QTableWidgetItem, QScrollArea)
from PyQt6.QtCore import Qt
import matplotlib.pyplot as plt
from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.backends.backend_qt5agg import NavigationToolbar2QT as NavigationToolbar
from openpyxl.styles import PatternFill
import matplotlib.dates as mdates

# 设置中文显示
plt.rcParams["font.family"] = ["SimHei", "Microsoft YaHei", "Arial Unicode MS"]
plt.rcParams['axes.unicode_minus'] = False  # 解决负号显示问题


class HoverTooltip:
    """用于实现鼠标悬停显示详细信息的工具类"""

    def __init__(self, fig, ax, data, x_col, y_col, tooltip_cols, is_categorical=False, hover_radius=None):
        self.fig = fig
        self.ax = ax
        self.data = data
        self.x_col = x_col
        self.y_col = y_col
        self.tooltip_cols = tooltip_cols
        self.is_categorical = is_categorical  # 标记x轴是否为分类数据

        # 可自定义悬停半径，None则使用默认值
        self.hover_radius = hover_radius

        # 存储点的位置和数据索引
        self.points = None
        self.annotations = []

        # 连接事件
        self.fig.canvas.mpl_connect('motion_notify_event', self.on_move)

    def set_points(self, points):
        """设置要监视的点集合"""
        self.points = points

    def on_move(self, event):
        """鼠标移动事件处理"""
        if event.inaxes != self.ax:
            self.hide_tooltips()
            return

        if self.points is None:
            return

        # 获取散点数据
        if hasattr(self.points, 'get_offsets'):
            offsets = self.points.get_offsets()
            x_coords = offsets[:, 0]
            y_coords = offsets[:, 1]
        else:
            x_coords = self.data[self.x_col]
            y_coords = self.data[self.y_col]

        # 检查鼠标是否在点的范围内
        hovered = False
        for i in range(len(x_coords)):
            # 处理不同类型的x轴数据
            if self.is_categorical:
                x = x_coords[i]
            elif isinstance(self.data.iloc[i][self.x_col], pd.Timestamp):
                x = mdates.date2num(self.data.iloc[i][self.x_col])
            else:
                x = x_coords[i]

            y = y_coords[i]

            # 计算合适的检测半径
            if self.hover_radius is not None:
                # 使用自定义半径
                threshold = self.hover_radius
            else:
                # 根据图表类型自动调整半径
                if self.is_categorical:
                    # 分类数据（箱线图）使用固定半径
                    threshold = 0.35  # 增大分类数据的检测半径
                else:
                    # 连续数据使用动态半径
                    x_range = self.ax.get_xlim()[1] - self.ax.get_xlim()[0]
                    y_range = self.ax.get_ylim()[1] - self.ax.get_ylim()[0]

                    # 对日期等大范围数据使用稍大比例，确保足够的检测范围
                    if x_range > 1000:  # 日期数据范围通常较大
                        threshold = max(x_range, y_range) * 0.015
                    else:
                        threshold = max(x_range, y_range) * 0.01

            # 检查鼠标是否在点的范围内
            if (abs(event.xdata - x) < threshold and
                    abs(event.ydata - y) < threshold):
                # 显示tooltip
                self.show_tooltip(i, event)
                hovered = True
                break

        if not hovered:
            self.hide_tooltips()

    def show_tooltip(self, index, event):
        """显示指定索引数据的详细信息"""
        self.hide_tooltips()

        # 构建tooltip文本
        tooltip_text = []
        for col in self.tooltip_cols:
            value = self.data.iloc[index][col]
            # 格式化日期
            if isinstance(value, pd.Timestamp):
                value = value.strftime('%Y-%m-%d')
            tooltip_text.append(f"{col}: {value}")

        # 创建注释
        annotation = self.ax.annotate(
            '\n'.join(tooltip_text),
            xy=(event.xdata, event.ydata),
            xytext=(10, 10),
            textcoords='offset points',
            bbox=dict(boxstyle='round,pad=0.5', fc='yellow', alpha=0.8),
            arrowprops=dict(arrowstyle='->', connectionstyle='arc3,rad=0')
        )

        self.annotations.append(annotation)
        self.fig.canvas.draw_idle()

    def hide_tooltips(self):
        """隐藏所有显示的tooltip"""
        for annotation in self.annotations:
            annotation.remove()
        self.annotations = []
        self.fig.canvas.draw_idle()


class ConcretePriceAnalyzer(QMainWindow):
    def __init__(self):
        super().__init__()
        self.data = None
        self.processed_data = None
        self.hover_toolips = []  # 存储所有悬停提示实例
        self.init_ui()

    def init_ui(self):
        # 设置窗口标题和大小
        self.setWindowTitle('混凝土价格分析工具')
        self.resize(1200, 800)

        # 创建主部件和布局
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        main_layout = QVBoxLayout(central_widget)

        # 创建菜单栏
        self.create_menu_bar()

        # 创建控制面板
        control_panel = self.create_control_panel()
        main_layout.addWidget(control_panel)

        # 创建标签页
        self.tabs = QTabWidget()
        self.create_tabs()
        main_layout.addWidget(self.tabs)

        self.show()

    def create_menu_bar(self):
        menubar = self.menuBar()

        # 文件菜单
        file_menu = menubar.addMenu('文件')

        # 加载数据动作
        load_action = file_menu.addAction('加载数据')
        load_action.triggered.connect(self.load_data)

        # 导出Excel动作
        export_action = file_menu.addAction('导出带异常标记的Excel')
        export_action.triggered.connect(self.export_excel)

        # 退出动作
        exit_action = file_menu.addAction('退出')
        exit_action.triggered.connect(self.close)

        # 帮助菜单
        help_menu = menubar.addMenu('帮助')
        about_action = help_menu.addAction('关于')
        about_action.triggered.connect(self.show_about)

    def create_control_panel(self):
        panel = QGroupBox("数据筛选")
        layout = QHBoxLayout()

        # 月份选择
        self.month_label = QLabel("选择月份:")
        self.month_combo = QComboBox()
        layout.addWidget(self.month_label)
        layout.addWidget(self.month_combo)

        # 规格选择
        self.spec_label = QLabel("选择规格:")
        self.spec_combo = QComboBox()
        layout.addWidget(self.spec_label)
        layout.addWidget(self.spec_combo)

        # 图表类型选择
        self.chart_label = QLabel("图表类型:")
        self.chart_combo = QComboBox()
        self.chart_combo.addItems(['箱线图', '散点图', '条形图(平均值)'])
        self.chart_combo.currentIndexChanged.connect(self.update_chart)
        layout.addWidget(self.chart_label)
        layout.addWidget(self.chart_combo)

        # 填充剩余空间
        layout.addStretch(1)

        panel.setLayout(layout)
        return panel

    def create_tabs(self):
        # 图表标签页
        self.chart_tab = QWidget()
        chart_layout = QVBoxLayout(self.chart_tab)

        # 创建Matplotlib图表
        self.figure = plt.figure(figsize=(10, 6))
        self.canvas = FigureCanvas(self.figure)
        self.toolbar = NavigationToolbar(self.canvas, self.chart_tab)

        chart_layout.addWidget(self.toolbar)
        chart_layout.addWidget(self.canvas)

        # 数据标签页
        self.data_tab = QWidget()
        data_layout = QVBoxLayout(self.data_tab)

        # 添加滚动区域以便查看大量数据
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        data_layout.addWidget(scroll_area)

        # 创建表格部件
        self.table_widget = QTableWidget()
        scroll_area.setWidget(self.table_widget)

        self.tabs.addTab(self.chart_tab, "图表分析")
        self.tabs.addTab(self.data_tab, "原始数据")

    def load_data(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "选择Excel文件", "", "Excel Files (*.xlsx *.xls)"
        )

        if file_path:
            try:
                self.data = self._load_and_preprocess_data(file_path)
                if self.data is not None:
                    self.processed_data = self._mark_outliers(self.data)
                    self._populate_combos()
                    self.update_chart()
                    self.update_data_table()  # 加载数据后更新表格
                    QMessageBox.information(self, "成功", "数据加载成功！")
            except Exception as e:
                QMessageBox.critical(self, "错误", f"加载数据失败: {str(e)}")

    def _load_and_preprocess_data(self, file_path):
        try:
            # 读取材料数据和项目映射表
            material_df = pd.read_excel(
                file_path,
                sheet_name="材料数据",
                engine='openpyxl'
            )
            project_df = pd.read_excel(
                file_path,
                sheet_name="项目映射表",
                engine='openpyxl'
            )

            # 数据预处理
            material_df['到货日期'] = pd.to_datetime(material_df['到货日期'])
            material_df['月份'] = material_df['到货日期'].dt.to_period('M')  # 提取月份
            material_df['规格'] = material_df['规格'].astype(str)  # 规格转为字符串

            # 合并地区信息
            merged_df = pd.merge(
                material_df,
                project_df[['项目名称', '地区（区/县）', '地区（市）']],
                on='项目名称',
                how='left'
            )

            return merged_df

        except Exception as e:
            print(f"数据加载错误: {e}")
            return None

    def _mark_outliers(self, data):
        # 标记异常值
        data_with_outliers = data.copy()
        grouped = data_with_outliers.groupby(['地区（市）', '规格'])
        data_with_outliers['is_outlier'] = False  # 初始化

        for (city, spec), group in grouped:
            outliers = self.detect_outliers_final(group, '单价（不含税）')
            data_with_outliers.loc[group.index, 'is_outlier'] = outliers

        return data_with_outliers

    def detect_outliers_final(self, data, column, diff_threshold=5, top_n=3, min_count=2):
        """异常值检测算法"""
        if data.empty or column not in data:
            return pd.Series([False] * len(data), index=data.index)

        # 提取目标列数据
        series = data[column].dropna()
        if series.empty:
            return pd.Series([False] * len(data), index=data.index)

        # 计算四分位数和箱线图上下限
        q1 = series.quantile(0.25)  # 25%分位数
        q3 = series.quantile(0.75)  # 75%分位数
        iqr = q3 - q1
        lower_bound = q1 - 1.5 * iqr  # 箱线图下限
        upper_bound = q3 + 1.5 * iqr  # 箱线图上限

        # 计算所有值的出现频率，按频率排序
        value_counts = series.value_counts().sort_values(ascending=False)

        # 筛选出符合最小出现次数的候选众数
        candidate_modes = [value for value, count in value_counts.items() if count >= min_count]

        # 取前N个众数，如果候选众数不足N个则取全部
        selected_modes = candidate_modes[:top_n]

        # 构建正常众数集合（众数及其近似值）
        normal_mode_values = set()
        for mode in selected_modes:
            normal_mode_values.add(mode)  # 添加众数本身
            # 添加与众数差值小于阈值的近似值
            for value in series.unique():
                if abs(value - mode) < diff_threshold and value not in normal_mode_values:
                    normal_mode_values.add(value)

        # 初始化异常值判断结果
        is_outlier = pd.Series([False] * len(series), index=series.index)

        # 遍历每个值，应用判断规则
        for idx, value in series.items():
            in_mode_range = value in normal_mode_values
            in_iqr_core = (value >= q1) and (value <= q3)  # 是否在25%-75%分位数之间
            in_box_range = (value >= lower_bound) and (value <= upper_bound)  # 是否在箱线图范围内

            # 应用四条规则
            if in_mode_range and in_box_range:
                is_outlier.loc[idx] = False  # 规则1: 众数内且在箱线内 → 正常
            elif in_mode_range and not in_box_range:
                is_outlier.loc[idx] = True  # 规则2: 众数内但超箱线 → 异常
            elif not in_mode_range and in_iqr_core:
                is_outlier.loc[idx] = False  # 规则3: 非众数但在25%-75%之间 → 正常
            else:  # not in_mode_range and not in_iqr_core
                is_outlier.loc[idx] = True  # 规则4: 非众数且不在25%-75%之间 → 异常

        return is_outlier.reindex(data.index, fill_value=False)

    def _populate_combos(self):
        # 填充月份下拉框
        months = sorted(self.processed_data['月份'].unique(), key=lambda x: str(x))
        self.month_combo.clear()
        for month in months:
            self.month_combo.addItem(month.strftime('%Y年%m月'), str(month))

        # 填充规格下拉框
        specs = sorted(self.processed_data['规格'].unique(),
                       key=lambda x: int(x[1:]) if x[1:].isdigit() else 0)
        self.spec_combo.clear()
        self.spec_combo.addItems(specs)

        # 绑定选择变化事件
        self.month_combo.currentIndexChanged.connect(self.update_chart)
        self.month_combo.currentIndexChanged.connect(self.update_data_table)  # 同步更新表格
        self.spec_combo.currentIndexChanged.connect(self.update_chart)
        self.spec_combo.currentIndexChanged.connect(self.update_data_table)  # 同步更新表格

    def update_chart(self):
        # 清除之前的悬停提示
        self.hover_toolips.clear()

        if self.processed_data is None:
            return

        # 获取当前选择
        selected_month = self.month_combo.currentData()
        selected_spec = self.spec_combo.currentText()
        chart_type = self.chart_combo.currentText()

        if not selected_month or not selected_spec:
            return

        # 筛选数据
        month_data = self.processed_data[
            (self.processed_data['月份'].astype(str) == selected_month) &
            (self.processed_data['规格'] == selected_spec)
            ].copy()

        month_name = pd.Period(selected_month).strftime('%Y年%m月')

        # 清除当前图表
        self.figure.clear()

        # 分离正常数据和异常数据
        normal_data = month_data[~month_data['is_outlier']]
        outlier_data = month_data[month_data['is_outlier']]

        # 创建图表
        ax = self.figure.add_subplot(111)

        if chart_type == '箱线图':
            # 获取所有唯一地区并排序，确保箱线图和散点图顺序一致
            regions = sorted(month_data['地区（市）'].unique())
            region_codes = {region: i for i, region in enumerate(regions)}

            # 为每个地区创建数据列表
            data_for_boxplot = [month_data[month_data['地区（市）'] == region]['单价（不含税）']
                                for region in regions]

            # 绘制箱线图
            bp = ax.boxplot(
                data_for_boxplot,
                patch_artist=True,  # 允许填充颜色
                widths=0.6  # 调整箱线图宽度
            )

            # 设置箱线图颜色
            for patch, color in zip(bp['boxes'], plt.cm.Set3.colors[:len(regions)]):
                patch.set_facecolor(color)

            # 设置x轴标签
            ax.set_xticks(range(1, len(regions) + 1))
            ax.set_xticklabels(regions, rotation=45, ha='right')

            # 添加随机噪声使同地区点分散显示，避免重叠
            jitter_strength = 0.1  # 抖动强度

            # 绘制正常数据点
            normal_x = [region_codes[region] + 1 + np.random.normal(0, jitter_strength)
                        for region in normal_data['地区（市）']]
            normal_points = ax.scatter(
                normal_x,
                normal_data['单价（不含税）'],
                color='black',
                alpha=0.5,
                s=50,
                label='正常数据'
            )

            # 绘制异常数据点
            outlier_points = None
            if not outlier_data.empty:
                outlier_x = [region_codes[region] + 1 + np.random.normal(0, jitter_strength)
                             for region in outlier_data['地区（市）']]
                outlier_points = ax.scatter(
                    outlier_x,
                    outlier_data['单价（不含税）'],
                    color='red',
                    marker='D',
                    s=80,
                    label='异常数据'
                )

            ax.set_xlabel('地区（市）')
            ax.set_ylabel('单价（不含税）')
            ax.legend()

            # 添加悬停提示 - 标记为分类数据，可自定义半径
            tooltip_cols = ['地区（市）', '项目名称', '到货日期', '单价（不含税）']
            if not normal_data.empty:
                normal_hover = HoverTooltip(
                    self.figure, ax, normal_data,
                    '地区（市）', '单价（不含税）', tooltip_cols,
                    is_categorical=True,
                    hover_radius=0.4  # 增大箱线图点的检测半径
                )
                normal_hover.set_points(normal_points)
                self.hover_toolips.append(normal_hover)

            if not outlier_data.empty and outlier_points is not None:
                outlier_hover = HoverTooltip(
                    self.figure, ax, outlier_data,
                    '地区（市）', '单价（不含税）', tooltip_cols,
                    is_categorical=True,
                    hover_radius=0.4  # 增大箱线图点的检测半径
                )
                outlier_hover.set_points(outlier_points)
                self.hover_toolips.append(outlier_hover)

        elif chart_type == '散点图':
            # 绘制正常数据
            normal_points = ax.scatter(
                normal_data['到货日期'],
                normal_data['单价（不含税）'],
                color='black',
                alpha=0.7,
                s=50,
                label='正常数据'
            )

            # 绘制异常数据
            outlier_points = None
            if not outlier_data.empty:
                outlier_points = ax.scatter(
                    outlier_data['到货日期'],
                    outlier_data['单价（不含税）'],
                    color='red',
                    marker='D',
                    s=80,
                    label='异常数据'
                )

            ax.set_xlabel('日期')
            ax.set_ylabel('单价（不含税）')
            ax.legend()
            # 旋转日期标签
            plt.xticks(rotation=45)

            # 添加悬停提示 - 为日期类散点图设置合适的半径
            tooltip_cols = ['地区（市）', '项目名称', '到货日期', '单价（不含税）']
            if not normal_data.empty:
                normal_hover = HoverTooltip(
                    self.figure, ax, normal_data,
                    '到货日期', '单价（不含税）', tooltip_cols,
                    hover_radius=None  # 使用自动计算的半径
                )
                normal_hover.set_points(normal_points)
                self.hover_toolips.append(normal_hover)

            if not outlier_data.empty and outlier_points is not None:
                outlier_hover = HoverTooltip(
                    self.figure, ax, outlier_data,
                    '到货日期', '单价（不含税）', tooltip_cols,
                    hover_radius=None  # 使用自动计算的半径
                )
                outlier_hover.set_points(outlier_points)
                self.hover_toolips.append(outlier_hover)

        elif chart_type == '条形图(平均值)':
            # 计算平均值和样本数量
            avg_data = month_data.groupby('地区（市）').agg(
                平均价格=('单价（不含税）', 'mean'),
                样本数量=('单价（不含税）', 'count')
            ).reset_index()

            # 按地区排序
            avg_data = avg_data.sort_values('地区（市）')

            # 绘制条形图
            bars = ax.bar(
                avg_data['地区（市）'],
                avg_data['平均价格'],
                color=plt.cm.skyblue(range(len(avg_data))),
                width=0.6
            )

            ax.set_xlabel('地区（市）')
            ax.set_ylabel('平均单价（不含税）')
            plt.xticks(rotation=45, ha='right')

            # 添加样本数量标签
            for bar, _, count in zip(bars, avg_data['平均价格'], avg_data['样本数量']):
                height = bar.get_height()
                ax.text(bar.get_x() + bar.get_width() / 2., height,
                        f"n={count}", ha='center', va='bottom')

        # 设置标题
        ax.set_title(f'{month_name} - 混凝土规格: {selected_spec}')

        # 调整布局
        self.figure.tight_layout()

        # 更新画布
        self.canvas.draw()

    def update_data_table(self):
        """更新原始数据表格显示"""
        if self.processed_data is None:
            return

        # 获取当前选择
        selected_month = self.month_combo.currentData()
        selected_spec = self.spec_combo.currentText()

        if not selected_month or not selected_spec:
            return

        # 筛选数据
        filtered_data = self.processed_data[
            (self.processed_data['月份'].astype(str) == selected_month) &
            (self.processed_data['规格'] == selected_spec)
            ].copy()

        # 选择要显示的列
        display_columns = ['地区（市）', '地区（区/县）', '项目名称', '到货日期',
                           '规格', '单价（不含税）', 'is_outlier']

        # 确保只包含存在的列
        display_columns = [col for col in display_columns if col in filtered_data.columns]
        table_data = filtered_data[display_columns].copy()

        # 格式化日期列
        for col in table_data.columns:
            if '日期' in col and pd.api.types.is_datetime64_any_dtype(table_data[col]):
                table_data[col] = table_data[col].dt.strftime('%Y-%m-%d')

        # 更新表格
        self.table_widget.setRowCount(len(table_data))
        self.table_widget.setColumnCount(len(display_columns))
        self.table_widget.setHorizontalHeaderLabels(display_columns)

        # 填充表格数据
        for row_idx, row_data in enumerate(table_data.itertuples(index=False)):
            for col_idx, value in enumerate(row_data):
                item = QTableWidgetItem(str(value))
                # 异常值行标红
                if display_columns[col_idx] == 'is_outlier' and value:
                    for c in range(len(display_columns)):
                        row_item = self.table_widget.item(row_idx, c)
                        if row_item:
                            row_item.setBackground(Qt.GlobalColor.red)
                        else:
                            item = QTableWidgetItem(str(table_data.iloc[row_idx, c]))
                            item.setBackground(Qt.GlobalColor.red)
                            self.table_widget.setItem(row_idx, c, item)
                elif display_columns[col_idx] != 'is_outlier':
                    self.table_widget.setItem(row_idx, col_idx, item)

        # 调整列宽
        self.table_widget.resizeColumnsToContents()

    def export_excel(self):
        if self.processed_data is None:
            QMessageBox.warning(self, "警告", "请先加载数据！")
            return

        output_file, _ = QFileDialog.getSaveFileName(
            self, "保存Excel文件", "混凝土价格数据_异常值标记.xlsx", "Excel Files (*.xlsx)"
        )

        if output_file:
            try:
                # 创建带异常标记的Excel
                df_with_outliers = self.processed_data.copy()

                # 写入Excel并标红异常值
                with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
                    df_with_outliers.to_excel(writer, sheet_name='价格数据（含异常标记）', index=False)
                    worksheet = writer.sheets['价格数据（含异常标记）']

                    # 红色填充样式
                    red_fill = PatternFill(start_color='FFFF0000', end_color='FFFF0000', fill_type='solid')

                    # 找到“是否异常”和“单价（不含税）”列的位置
                    is_outlier_col = None
                    price_col = None
                    for col in range(worksheet.max_column):
                        if worksheet.cell(row=1, column=col + 1).value == 'is_outlier':
                            is_outlier_col = col + 1
                        if worksheet.cell(row=1, column=col + 1).value == '单价（不含税）':
                            price_col = col + 1

                    # 标红异常值
                    if is_outlier_col and price_col:
                        for row in range(2, worksheet.max_row + 1):
                            if worksheet.cell(row=row, column=is_outlier_col).value is True:
                                worksheet.cell(row=row, column=price_col).fill = red_fill

                QMessageBox.information(self, "成功", f"已生成带异常值标记的Excel文件: {output_file}")
            except Exception as e:
                QMessageBox.critical(self, "错误", f"导出Excel失败: {str(e)}")

    def show_about(self):
        QMessageBox.about(self, "关于", "混凝土价格分析工具\n版本 1.0\n用于分析混凝土价格数据并检测异常值")


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = ConcretePriceAnalyzer()
    sys.exit(app.exec())
